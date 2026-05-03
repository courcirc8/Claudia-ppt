"""Accessibility & layout audit (P5).

Checks:
- Image bounds (off-slide = ERROR; margin <0.5" = WARNING)
- Font size minimums (body <15pt, title <24pt = WARNING)
- Contrast ratio between text color and (estimated) bg (WCAG AA = 4.5:1)
- Alt-text presence on images
- One title per slide hierarchy
"""
from __future__ import annotations
from dataclasses import dataclass
from pptx import Presentation


@dataclass
class Issue:
    slide: int
    shape: int | None
    severity: str  # "ERROR" | "WARNING" | "INFO"
    code: str
    message: str
    fix: str | None = None


def _luminance(rgb: tuple[int, int, int]) -> float:
    def channel(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (channel(x) for x in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast_ratio(fg: tuple[int, int, int], bg: tuple[int, int, int]) -> float:
    l1, l2 = _luminance(fg), _luminance(bg)
    lighter, darker = max(l1, l2), min(l1, l2)
    return (lighter + 0.05) / (darker + 0.05)


def _slide_bg_rgb(slide) -> tuple[int, int, int] | None:
    """Best-effort: walk the slide bg fill or default white."""
    try:
        bg = slide.background
        fill = bg.fill
        if fill.type is not None and hasattr(fill, "fore_color"):
            try:
                rgb = fill.fore_color.rgb
                return (rgb[0], rgb[1], rgb[2])
            except (AttributeError, KeyError):
                pass
    except Exception:
        pass
    return (255, 255, 255)  # default white


def _shape_text_color(shape) -> tuple[int, int, int] | None:
    if not shape.has_text_frame:
        return None
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            try:
                rgb = run.font.color.rgb
                return (rgb[0], rgb[1], rgb[2])
            except Exception:
                continue
    return None


def _first_font_size(shape) -> int | None:
    if not shape.has_text_frame:
        return None
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                return int(run.font.size.pt)
    return None


def audit_accessibility(
    pptx_path: str,
    *,
    min_body_pt: int = 15,
    min_title_pt: int = 24,
    min_margin_in: float = 0.5,
    wcag_aa_ratio: float = 4.5,
) -> list[Issue]:
    """Run full a11y + layout audit. Returns sorted Issue list (errors first)."""
    p = Presentation(pptx_path)
    SW, SH = p.slide_width / 914400, p.slide_height / 914400
    issues: list[Issue] = []

    for si, slide in enumerate(p.slides):
        bg_rgb = _slide_bg_rgb(slide) or (255, 255, 255)
        title_count = 0

        for shi, shape in enumerate(slide.shapes):
            L = shape.left / 914400
            T = shape.top / 914400
            W = shape.width / 914400
            H = shape.height / 914400
            R = L + W
            B = T + H

            # 1. Bounds (any shape)
            if R > SW + 0.01 or B > SH + 0.01:
                issues.append(Issue(si, shi, "ERROR", "out_of_bounds",
                    f"shape exits slide: right={R:.2f}>{SW:.2f} or bottom={B:.2f}>{SH:.2f}",
                    fix=f"Move/resize so right≤{SW:.2f} and bottom≤{SH:.2f}"))

            # 2. Image margins (cover/closing-style)
            if shape.shape_type == 13:  # picture
                if L > 0.1 and (SW - R) < min_margin_in:
                    issues.append(Issue(si, shi, "WARNING", "tight_image_margin",
                        f"image right margin {SW-R:.2f}\" < {min_margin_in}\"",
                        fix="Move image left or shrink width"))

            # 3. Font size
            if shape.has_text_frame:
                fs = _first_font_size(shape)
                if fs is not None:
                    is_title = T < 1.3
                    minimum = min_title_pt if is_title else min_body_pt
                    if fs < minimum:
                        issues.append(Issue(si, shi, "WARNING", "font_too_small",
                            f"{'title' if is_title else 'body'} font {fs}pt < min {minimum}pt",
                            fix=f"Increase to {minimum}pt or larger"))
                    if is_title:
                        title_count += 1

                # 4. Contrast
                fg = _shape_text_color(shape)
                if fg:
                    ratio = contrast_ratio(fg, bg_rgb)
                    if ratio < wcag_aa_ratio:
                        issues.append(Issue(si, shi, "WARNING", "low_contrast",
                            f"contrast {ratio:.1f}:1 < WCAG AA {wcag_aa_ratio}:1 (text {fg} on bg {bg_rgb})",
                            fix="Darken text or lighten bg"))

            # 5. Alt-text on pictures
            if shape.shape_type == 13:
                desc = (shape.element.xpath(".//@descr") or [""])[0]
                if not desc:
                    issues.append(Issue(si, shi, "INFO", "missing_alt_text",
                        "image has no alt-text",
                        fix="Set shape.descr for screen readers"))

        # 6. One title per slide
        if title_count == 0:
            issues.append(Issue(si, None, "INFO", "no_title",
                "no title-positioned text found", fix="Add a top-positioned title text box"))
        elif title_count > 1:
            issues.append(Issue(si, None, "INFO", "multiple_titles",
                f"{title_count} title-positioned text boxes",
                fix="Consider consolidating to one H1 per slide"))

    sev_order = {"ERROR": 0, "WARNING": 1, "INFO": 2}
    issues.sort(key=lambda i: (sev_order[i.severity], i.slide, i.shape or -1))
    return issues


def format_issues(issues: list[Issue]) -> str:
    if not issues:
        return "✓ no issues found"
    lines = []
    for iss in issues:
        loc = f"slide {iss.slide}" + (f" shape {iss.shape}" if iss.shape is not None else "")
        lines.append(f"[{iss.severity}] {loc} — {iss.code}: {iss.message}"
                    + (f"\n    → {iss.fix}" if iss.fix else ""))
    return "\n".join(lines)
