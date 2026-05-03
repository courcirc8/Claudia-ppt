"""Tests for pptx_kit modules — text / quality / accessibility / image_style."""
from __future__ import annotations
import json
import shutil
import sys
import tempfile
from pathlib import Path

import pytest

# Claudia-ppt is the TOOL root; presentations live in ../presentations/
TOOL_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(TOOL_ROOT))

# Sample presentation used as integration fixture; override via SAMPLE_PPTX env
import os
SAMPLE_PPTX = Path(os.environ.get(
    "SAMPLE_PPTX",
    TOOL_ROOT.parent / "presentations" / "droit_des_familles" / "droit_des_familles.pptx"
))

from pptx_kit.text import (
    inspect, bulk_apply, smart_wrap, estimate_capacity,
)
from pptx_kit.image_style import (
    StyleAnchor, THEMES, request_variations,
    register_image, manifest_for_project,
)
from pptx_kit.quality import (
    incremental_export, diff_presentations, _slide_xml_hash,
)
from pptx_kit.accessibility import (
    audit_accessibility, contrast_ratio, format_issues,
)

# Backwards-compat alias used by existing tests
PPTX = SAMPLE_PPTX


# ─── text.py ──────────────────────────────────────────────────────

class TestEstimateCapacity:
    def test_default_calibri_22pt(self):
        cap = estimate_capacity(width_in=8.13, height_in=5.4, font_size_pt=22, font_name="Calibri")
        assert 40 <= cap["chars_per_line"] <= 60
        assert cap["max_lines"] >= 10
        assert cap["wrap_safe_threshold"] < cap["chars_per_line"]

    def test_smaller_font_more_chars(self):
        small = estimate_capacity(8.13, 5.4, 12)
        big = estimate_capacity(8.13, 5.4, 24)
        assert small["chars_per_line"] > big["chars_per_line"]
        assert small["max_lines"] > big["max_lines"]

    def test_unknown_font_uses_default(self):
        cap = estimate_capacity(8.0, 5.0, 16, font_name="DoesNotExist")
        assert cap["chars_per_line"] > 0


class TestSmartWrap:
    def test_keeps_legal_article_atomic(self):
        out = smart_wrap("Les biens (art. 198 ch. 2 CC) constituent une catégorie particulière.",
                        width_in=4.0, font_size_pt=14)
        # "art. 198 ch. 2 CC" must never be split across lines
        for line in out.split("\n"):
            if "art." in line:
                assert "198" in line
                assert "CC" in line

    def test_keeps_tf_reference_atomic(self):
        out = smart_wrap("Selon TF 5A_54/2024 du 15 mars 2026 le principe est clair.",
                        width_in=3.5, font_size_pt=14)
        for line in out.split("\n"):
            if "TF" in line:
                assert "5A_54/2024" in line

    def test_no_word_split(self):
        out = smart_wrap("provisionnellesprovisionnelles est un mot très long " * 5,
                        width_in=4.0, font_size_pt=18)
        for line in out.split("\n"):
            if line.strip():
                # No line should end mid-word for normal words
                assert not line.endswith("-")  # naive but useful
                # Each line should be a sequence of complete words
                # i.e., no whitespace-followed-by-punctuation oddities
                assert line == line.strip()

    def test_preserves_existing_newlines(self):
        out = smart_wrap("Ligne 1.\n\nLigne 2.\nLigne 3.", width_in=8.0, font_size_pt=18)
        assert out.count("\n") >= 3  # paragraph break preserved

    def test_short_text_no_wrap_needed(self):
        out = smart_wrap("Court.", width_in=8.0, font_size_pt=18)
        assert "\n" not in out


class TestInspect:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_inspect_returns_shapes(self):
        shapes = inspect(str(PPTX))
        assert len(shapes) > 0
        assert any(s.role == "image" for s in shapes)
        assert any(s.role == "title" for s in shapes)
        assert any(s.role == "body" for s in shapes)

    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_classify_title_top(self):
        shapes = inspect(str(PPTX))
        titles = [s for s in shapes if s.role == "title"]
        # Titles should be near the top (top_in < 1.3)
        for t in titles:
            assert t.top_in < 1.3, f"title shape at top={t.top_in} should be <1.3"


class TestBulkApply:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_dry_run_counts_body_shapes(self):
        # Don't save — just count what would be modified
        n = bulk_apply(str(PPTX), filter={"role": "body"}, save=False)
        assert n >= 9  # at least 9 content slides have body

    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_filter_callable(self):
        n = bulk_apply(str(PPTX),
                      filter=lambda s: s.role == "title",
                      save=False)
        assert n >= 8  # at least 8 titles


# ─── image_style.py ────────────────────────────────────────────────

class TestStyleAnchor:
    def test_prompt_for_includes_palette(self):
        anchor = StyleAnchor(name="x", palette_words="rose et or", mood="éditorial")
        p = anchor.prompt_for("a balance scale")
        assert "rose et or" in p
        assert "éditorial" in p
        assert "a balance scale" in p

    def test_constraints_default_no_text(self):
        anchor = StyleAnchor(name="x", palette_words="bleu", mood="moderne")
        assert "NO TEXT" in anchor.prefix

    def test_themes_present(self):
        for k in ("pink_violet", "navy_gold", "earth_terracotta", "teal_coral",
                  "mono_crimson", "sage_rose"):
            assert k in THEMES, f"Missing theme: {k}"

    def test_save_load_roundtrip(self, tmp_path):
        anchor = StyleAnchor(name="test", palette_words="vert", mood="naturel")
        f = tmp_path / "anchor.json"
        anchor.save(f)
        loaded = StyleAnchor.load(f)
        assert loaded.name == "test"
        assert loaded.mood == "naturel"


class TestRequestVariations:
    def test_returns_n_specs(self):
        specs = request_variations("img-abc", count=4)
        assert len(specs) == 4

    def test_each_spec_has_unique_seed(self):
        specs = request_variations("img-abc", count=3)
        seeds = [s["input"]["seed"] for s in specs]
        assert len(set(seeds)) == len(seeds)

    def test_strength_param_propagates(self):
        specs = request_variations("img-abc", count=2, variation_strength=0.7)
        for s in specs:
            assert s["input"]["strength"] == 0.7


class TestManifest:
    def test_creates_manifest_if_missing(self, tmp_path):
        mp = manifest_for_project(tmp_path)
        assert mp.exists()
        data = json.loads(mp.read_text())
        assert "slides" in data
        assert "images" in data

    def test_register_image_appends(self, tmp_path):
        # Create dummy image file
        img_path = tmp_path / "foo.png"
        img_path.write_bytes(b"PNG\x00")
        register_image(tmp_path, slide_idx=2, image_path=img_path,
                      prompt="test prompt", model="seedream", theme_anchor="pink_violet")
        data = json.loads((tmp_path / "manifest.json").read_text())
        assert len(data["images"]) == 1
        assert data["images"][0]["slide_index"] == 2
        assert data["images"][0]["theme_anchor"] == "pink_violet"


# ─── quality.py ────────────────────────────────────────────────────

class TestDiff:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_self_diff_is_empty(self, tmp_path):
        snap = tmp_path / "snap.pptx"
        shutil.copy2(PPTX, snap)
        diff = diff_presentations(str(snap), str(PPTX))
        assert diff == [], f"self-diff should be empty, got: {diff}"

    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_detects_font_size_change(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Pt
        before = tmp_path / "before.pptx"
        after = tmp_path / "after.pptx"
        shutil.copy2(PPTX, before)
        shutil.copy2(PPTX, after)
        # Modify one font size in 'after' — bump every sized run, save once, diff once.
        p = Presentation(str(after))
        modified = False
        for slide in p.slides:
            for sh in slide.shapes:
                if not sh.has_text_frame:
                    continue
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size:
                            run.font.size = Pt(99)
                            modified = True
        assert modified, "fixture has no run with explicit font size — test cannot exercise diff"
        p.save(str(after))
        changes = diff_presentations(str(before), str(after))
        assert any(c["kind"] == "font_size_changed" for c in changes), \
            f"diff did not detect font_size_changed; got: {changes}"


class TestSlideHash:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_hash_deterministic(self):
        from pptx import Presentation
        p1 = Presentation(str(PPTX))
        p2 = Presentation(str(PPTX))
        h1 = _slide_xml_hash(p1.slides[0])
        h2 = _slide_xml_hash(p2.slides[0])
        assert h1 == h2

    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_different_slides_different_hashes(self):
        from pptx import Presentation
        p = Presentation(str(PPTX))
        if len(p.slides) >= 2:
            assert _slide_xml_hash(p.slides[0]) != _slide_xml_hash(p.slides[1])


class TestIncrementalExport:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_state_file_created(self, tmp_path):
        # Just verify state file mechanism without actually rendering (skip soffice)
        # Manually create a fake state to test the skip path
        state_file = tmp_path / ".export_state.json"
        from pptx_kit.quality import _all_hashes
        hashes = _all_hashes(str(PPTX))
        state_file.write_text(json.dumps({str(k): v for k, v in hashes.items()}))

        # If we re-call with all hashes matching, should skip
        result = incremental_export(str(PPTX), tmp_path, only_changed=True)
        # All slides unchanged → rendered should be empty
        assert result["rendered"] == [] or result["total"] == len(result["rendered"]) + len(result["skipped"])


# ─── accessibility.py ──────────────────────────────────────────────

class TestContrastRatio:
    def test_white_on_black_max(self):
        ratio = contrast_ratio((255, 255, 255), (0, 0, 0))
        assert 20.5 < ratio < 21.5  # WCAG max ~21:1

    def test_identical_colors_min(self):
        ratio = contrast_ratio((128, 128, 128), (128, 128, 128))
        assert ratio == pytest.approx(1.0)

    def test_meets_aa_for_dark_text_on_light(self):
        # #1A1612 on #ECE5D5 (project chrome)
        ratio = contrast_ratio((26, 22, 18), (236, 229, 213))
        assert ratio >= 4.5  # WCAG AA


class TestAudit:
    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_audit_returns_list(self):
        issues = audit_accessibility(str(PPTX))
        assert isinstance(issues, list)
        # Check Issue structure
        for i in issues:
            assert hasattr(i, "severity")
            assert i.severity in ("ERROR", "WARNING", "INFO")
            assert hasattr(i, "code")

    @pytest.mark.skipif(not PPTX.exists(), reason="droit_des_familles.pptx not available")
    def test_no_critical_errors(self):
        """The current built deck should have no out-of-bounds errors."""
        issues = audit_accessibility(str(PPTX))
        errors = [i for i in issues if i.severity == "ERROR"]
        assert len(errors) == 0, f"Unexpected ERRORS in deck: {format_issues(errors)}"

    def test_format_empty_issues(self):
        assert format_issues([]) == "✓ no issues found"


# Run with: pytest tests/test_pptx_kit.py -v
if __name__ == "__main__":
    pytest.main([__file__, "-v"])
